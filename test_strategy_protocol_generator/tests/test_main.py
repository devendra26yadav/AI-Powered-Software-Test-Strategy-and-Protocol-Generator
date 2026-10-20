from io import BytesIO
import json
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from test_strategy_protocol_generator.main import (
    DocumentationMode,
    GENERAL_TEST_DESIGN_OUTPUT_INSTRUCTIONS,
    GENERAL_REQUIREMENTS_OUTPUT_INSTRUCTIONS,
    GENERAL_SOFTWARE_DOCUMENTS,
    DOCUMENT_MODE_PROFILES,
    MEDICAL_DEVICE_MODE_INSTRUCTIONS,
    MEDICAL_DEVICE_PROTOCOL_INSTRUCTIONS,
    MEDICAL_DEVICE_SOFTWARE_DOCUMENTS,
    MEDICAL_DEVICE_SUMMARY_INSTRUCTIONS,
    MEDICAL_DEVICE_TEST_PLAN_INSTRUCTIONS,
    MEDICAL_DEVICE_TRACEABILITY_INSTRUCTIONS,
    _build_crew_inputs,
    _build_document_contents,
    _build_traceability_matrix,
    _extract_markdown_section,
    _has_verification_execution_evidence,
    _mark_cases_not_executed,
    _remove_markdown_section,
    _run_generation_pipeline,
    _validate_document_quality,
    _write_document_package,
)


def _read_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    blocks.extend(
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    )
    return "\n".join(blocks)


def _read_xlsx(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active
    return "\n".join(
        " | ".join(str(value or "") for value in row)
        for row in worksheet.iter_rows(values_only=True)
    )


def _zip_members(content: bytes) -> list[str]:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        return archive.namelist()


class TaskOutput:
    def __init__(self, raw):
        self.raw = raw


class BuildCrewInputsTests(unittest.TestCase):
    def test_defaults_to_general_software_mode(self):
        inputs = _build_crew_inputs("Example project", "Project context")

        self.assertEqual(inputs["documentation_mode"], DocumentationMode.GENERAL_SOFTWARE.value)
        self.assertEqual(inputs["topic"], "Example project")
        self.assertEqual(inputs["document_context"], "Project context")
        self.assertIn(
            "Requirements Traceability Matrix",
            inputs["test_design_output_instructions"],
        )
        self.assertEqual(
            inputs["requirements_output_instructions"],
            GENERAL_REQUIREMENTS_OUTPUT_INSTRUCTIONS,
        )
        self.assertIn("REQ-F-001", inputs["requirements_output_instructions"])

    def test_preserves_selected_medical_device_mode(self):
        inputs = _build_crew_inputs(
            "Example device",
            "Device context",
            DocumentationMode.MEDICAL_DEVICE_SOFTWARE.value,
        )

        self.assertEqual(
            inputs["documentation_mode"],
            DocumentationMode.MEDICAL_DEVICE_SOFTWARE.value,
        )
        self.assertEqual(
            inputs["test_design_output_instructions"],
            MEDICAL_DEVICE_TRACEABILITY_INSTRUCTIONS,
        )
        self.assertEqual(
            inputs["medical_device_mode_instructions"],
            MEDICAL_DEVICE_MODE_INSTRUCTIONS,
        )
        self.assertEqual(
            inputs["medical_device_test_plan_instructions"],
            MEDICAL_DEVICE_TEST_PLAN_INSTRUCTIONS,
        )
        self.assertEqual(
            inputs["medical_device_protocol_instructions"],
            MEDICAL_DEVICE_PROTOCOL_INSTRUCTIONS,
        )
        self.assertTrue(
            inputs["medical_device_summary_instructions"].startswith(
                MEDICAL_DEVICE_SUMMARY_INSTRUCTIONS
            )
        )
        self.assertIn(
            "Pending / Not Executed",
            inputs["medical_device_summary_instructions"],
        )
        self.assertIn("IEC 62304", inputs["medical_device_mode_instructions"])
        self.assertIn("ISO 14971", inputs["medical_device_mode_instructions"])
        self.assertIn("ISO 13485", inputs["medical_device_mode_instructions"])
        self.assertIn("FDA QMSR", inputs["medical_device_mode_instructions"])
        self.assertIn("Risk Control", inputs["test_design_output_instructions"])
        self.assertIn("REQ-F-001", inputs["requirements_output_instructions"])

    def test_extract_markdown_section_stops_at_next_peer_heading(self):
        content = (
            "## Requirements Traceability Matrix\nMatrix content\n"
            "## Detailed Test Cases\nCase content"
        )

        section = _extract_markdown_section(content, "Requirements Traceability Matrix")

        self.assertEqual(section, "## Requirements Traceability Matrix\nMatrix content")

    def test_remove_markdown_section_preserves_other_sections(self):
        content = (
            "## Requirements Traceability Matrix\nMatrix content\n"
            "## Test Cases\nCase content"
        )

        remaining = _remove_markdown_section(content, "Requirements Traceability Matrix")

        self.assertEqual(remaining, "## Test Cases\nCase content")

    def test_execution_evidence_requires_case_result_status_and_reference(self):
        context = (
            "Test Case ID: TC-1\n"
            "Actual Results: Inventory count updated\n"
            "Test Status: Passed\n"
            "Objective Evidence Reference: run-123/screenshot.png"
        )

        self.assertTrue(_has_verification_execution_evidence(context))
        self.assertFalse(
            _has_verification_execution_evidence(
                context.replace("Objective Evidence Reference: run-123/screenshot.png", "")
            )
        )

    def test_traceability_adds_pending_rows_for_unmapped_ids(self):
        from test_strategy_protocol_generator.main import DOCUMENT_MODE_PROFILES

        matrix = _build_traceability_matrix(
            "REQ-F-001: Functional requirement",
            "RISK-001; RC-001",
            "## Requirements and Risk Traceability Matrix\n"
            "| Requirement ID | Risk Reference | Risk Control | Test Case ID | "
            "Verification Status | Evidence Reference |\n"
            "|---|---|---|---|---|---|\n"
            "| REQ-F-001 | Not provided | Not provided | TC-1 | Pending / Not Verified | Not provided |",
            DOCUMENT_MODE_PROFILES[DocumentationMode.MEDICAL_DEVICE_SOFTWARE],
        )

        self.assertIn(
            "| Not provided | RISK-001 | Not provided | Not linked | "
            "Pending / Not Verified | Not provided |",
            matrix,
        )
        self.assertIn(
            "| Not provided | Not provided | RC-001 | Not linked | "
            "Pending / Not Verified | Not provided |",
            matrix,
        )
        self.assertIn("Risk identifiers requiring mapping: RISK-001, RC-001", matrix)

    def test_traceability_preserves_rsk_risk_identifiers(self):
        matrix = _build_traceability_matrix(
            "REQ-F-001: Functional requirement",
            "RSK-005: Unauthorized stock change",
            "## Requirements and Risk Traceability Matrix\n"
            "| Requirement ID | Risk Reference | Risk Control | Test Case ID | "
            "Verification Status | Evidence Reference |\n"
            "|---|---|---|---|---|---|\n"
            "| REQ-F-001 | RSK-005 | Authorized adjustment control | TC-1 | "
            "Pending / Not Verified | Not provided |",
            DOCUMENT_MODE_PROFILES[DocumentationMode.MEDICAL_DEVICE_SOFTWARE],
        )

        self.assertIn("RSK-005", matrix)
        self.assertNotIn("Risk identifiers requiring mapping: RSK-005", matrix)

    def test_case_execution_fields_are_cleared_without_execution_evidence(self):
        cases = (
            "- **Actual Results:** It worked\n"
            "- **Test Status:** Passed\n"
            "- **Objective Evidence Reference:** evidence-1"
        )

        pending = _mark_cases_not_executed(cases)

        self.assertIn("Actual Results:** Not executed", pending)
        self.assertIn("Test Status:** Not Executed", pending)
        self.assertIn("Objective Evidence Reference:** Not provided", pending)

    def test_protocol_case_details_are_not_duplicated_from_protocol_task(self):
        outputs = [
            TaskOutput("Test plan"),
            TaskOutput("Research"),
            TaskOutput("REQ-F-001: Requirement"),
            TaskOutput("Strategy"),
            TaskOutput("Risk"),
            TaskOutput(
                "Protocol execution guidance.\n\n"
                "Test Case ID: TC-1\n"
                "Linked Requirement ID: REQ-F-001\n"
                "Expected Results: Duplicate generated case."
            ),
            TaskOutput(
                "## Requirements Traceability Matrix\n"
                "| Requirement ID | Test Case ID | Verification Status | Evidence Reference |\n"
                "|---|---|---|---|\n"
                "| REQ-F-001 | TC-1 | Pending | Not provided |\n\n"
                "## Detailed Test Cases\n"
                "### TC-1\n"
                "- **Linked Requirement ID:** REQ-F-001\n"
                "- **Expected Results:** Canonical generated case."
            ),
            TaskOutput("Automation"),
            TaskOutput("Quality"),
            TaskOutput("Report"),
        ]

        protocol = _build_document_contents(
            outputs, type("CrewResult", (), {"raw": ""})(), DocumentationMode.GENERAL_SOFTWARE
        )[1]

        self.assertIn("Protocol execution guidance.", protocol)
        self.assertIn("Canonical generated case.", protocol)
        self.assertNotIn("Duplicate generated case.", protocol)
        self.assertEqual(protocol.count("TC-1"), 1)

    def test_shared_pipeline_selects_profile_for_each_mode(self):
        for mode, documents in (
            (DocumentationMode.GENERAL_SOFTWARE, GENERAL_SOFTWARE_DOCUMENTS),
            (DocumentationMode.MEDICAL_DEVICE_SOFTWARE, MEDICAL_DEVICE_SOFTWARE_DOCUMENTS),
        ):
            with self.subTest(mode=mode.value):
                outputs = [TaskOutput("") for _ in range(10)]
                built = _build_document_contents(
                    outputs,
                    type("CrewResult", (), {"raw": "Summary"})(),
                    mode,
                )
                self.assertEqual(len(built), len(documents))

    def test_shared_generation_pipeline_passes_mode_to_crew_and_export(self):
        result = type("CrewResult", (), {"tasks_output": [], "raw": "Summary"})()
        for mode in DocumentationMode:
            with self.subTest(mode=mode.value):
                with (
                    patch("test_strategy_protocol_generator.main.TestStrategyProtocolGenerator") as crew_class,
                    patch("test_strategy_protocol_generator.main._write_document_package") as writer,
                ):
                    crew_class.return_value.crew.return_value.kickoff.return_value = result
                    writer.return_value = (
                        ["document.md"],
                        "documents.zip",
                        {"status": "PASS", "checks": []},
                    )

                    returned_result, files, archive, status, report = _run_generation_pipeline(
                        " Example ", [], mode
                    )

                    self.assertIs(returned_result, result)
                    self.assertEqual((files, archive), (["document.md"], "documents.zip"))
                    self.assertIn("Generated 4", status)
                    self.assertIn("PASS", status)
                    self.assertEqual(report["status"], "PASS")
                    kickoff_inputs = crew_class.return_value.crew.return_value.kickoff.call_args.kwargs["inputs"]
                    self.assertEqual(kickoff_inputs["documentation_mode"], mode.value)
                    writer.assert_called_once()
                    self.assertEqual(writer.call_args.args[2], mode)


class DocumentPackageTests(unittest.TestCase):
    def setUp(self):
        self.task_outputs = [
            TaskOutput("Project scope and acceptance criteria"),
            TaskOutput("Research findings"),
            TaskOutput(
                "Functional requirements:\n"
                "- REQ-F-001: Inventory items can be created.\n"
                "Nonfunctional requirements:\n"
                "- REQ-NF-001: Standard pages load within two seconds."
            ),
            TaskOutput("Test strategy and environment"),
            TaskOutput("RISK-001 with RC-001 risk control"),
            TaskOutput("Test execution procedures"),
            TaskOutput(
                "## Requirements Traceability Matrix\n"
                "| Requirement ID | Test Case ID | Verification Status | Evidence Reference |\n"
                "| REQ-F-001 | TC-1 | Pending | Not provided |\n\n"
                "## Detailed Test Cases\n"
                "### TC-1\n"
                "TC-1: Verify inventory item creation.\n"
                "- Requirement ID: REQ-F-001"
            ),
            TaskOutput("Manual and automated testing approach"),
            TaskOutput("Quality review"),
            TaskOutput("Executive conclusions"),
        ]

    def test_quality_report_flags_all_required_validation_categories(self):
        mode = DocumentationMode.MEDICAL_DEVICE_SOFTWARE
        documents = {
            filename: (
                f"# {title}\n\nProject: Example project\n\n"
                f"Documentation Mode: {mode.value}\n\n"
            )
            for filename, title in MEDICAL_DEVICE_SOFTWARE_DOCUMENTS
        }
        documents["Software_Verification_Protocol.docx"] += (
            "### TC-1\n"
            "- **Linked Requirement ID:** REQ-F-999\n"
            "- **Linked Risk Control ID:** RC-404\n"
            "- **Objective:** Save an item\n"
            "- **Test Steps:** Save an item\n"
            "- **Expected Results:** Item is saved\n"
            "- **Actual Results:** Item saved\n"
            "- **Test Status:** Passed\n"
            "- **Objective Evidence Reference:** fake-run\n\n"
            "### TC-2\n"
            "- **Linked Requirement ID:** REQ-F-001\n"
            "- **Linked Risk Control ID:** RC-404\n"
            "- **Objective:** Save an item\n"
            "- **Test Steps:** Save an item\n"
            "- **Expected Results:** Item is saved\n\n"
            "### TC-1\n"
            "- **Linked Requirement ID:** REQ-F-001\n"
            "- **Linked Risk Control ID:** RC-404\n"
            "- **Objective:** Save an item\n"
            "- **Test Steps:** Save an item\n"
            "- **Expected Results:** Item is overwritten\n"
        )
        documents["Requirements_Risk_Traceability_Matrix.xlsx"] += (
            "| Requirement ID | Risk Reference | Risk Control | Test Case ID | "
            "Verification Status | Evidence Reference |\n"
            "|---|---|---|---|---|---|\n"
            "| REQ-F-001 | RISK-1 | RC-404 | TC-404 | Pending | Not provided |\n"
        )
        documents["Software_Verification_Summary_Report.docx"] += (
            "### TC-1\n- **Expected Results:** A different result\n"
        )
        documents["Software_Test_Plan.docx"] = documents[
            "Software_Test_Plan.docx"
        ].replace("Project: Example project", "Project: Different project")

        report = _validate_document_quality(
            "Example project",
            mode,
            documents,
            "REQ-F-001: Create item\nREQ-NF-001: Load in two seconds",
            "RISK-1 with control RC-1",
            "",
            ["TC-1: unsupported execution claim removed."],
        )

        self.assertEqual(len(report["checks"]), 9)
        self.assertEqual(report["status"], "REVIEW_REQUIRED")
        check_status = {item["id"]: item["status"] for item in report["checks"]}
        self.assertEqual(check_status["unique_test_case_identifiers"], "FAIL")
        self.assertEqual(check_status["valid_requirement_references"], "FAIL")
        self.assertEqual(check_status["traceability_matrix_test_cases"], "FAIL")
        self.assertEqual(check_status["expectation_consistency"], "FAIL")
        self.assertEqual(check_status["duplicate_test_cases"], "FAIL")
        self.assertEqual(check_status["requirements_coverage"], "REVIEW")
        self.assertEqual(check_status["execution_results_supported"], "REVIEW")
        self.assertEqual(check_status["medical_risk_control_traceability"], "FAIL")
        self.assertEqual(check_status["document_metadata_consistency"], "FAIL")

    def test_valid_document_set_passes_quality_validation(self):
        mode = DocumentationMode.GENERAL_SOFTWARE
        documents = {
            filename: (
                f"# {title}\n\nProject: Example project\n\n"
                f"Documentation Mode: {mode.value}\n\n"
            )
            for filename, title in GENERAL_SOFTWARE_DOCUMENTS
        }
        documents["Software_Test_Protocol.docx"] += (
            "### TC-1\n"
            "- **Linked Requirement ID:** REQ-F-001\n"
            "- **Test Steps:** Create an item\n"
            "- **Expected Results:** Item is created\n"
        )
        documents["Requirements_Traceability_Matrix.xlsx"] += (
            "| Requirement ID | Test Case ID | Verification Status | Evidence Reference |\n"
            "|---|---|---|---|\n"
            "| REQ-F-001 | TC-1 | Pending | Not provided |\n"
        )

        report = _validate_document_quality(
            "Example project",
            mode,
            documents,
            "REQ-F-001: Create item",
            "",
            "",
        )

        self.assertEqual(report["status"], "PASS")
        self.assertTrue(all(check["status"] == "PASS" for check in report["checks"]))

    def _write_package(
        self,
        mode=DocumentationMode.GENERAL_SOFTWARE,
        task_outputs=None,
        document_context="",
    ):
        return _write_document_package(
            type(
                "CrewResult",
                (),
                {
                    "tasks_output": task_outputs or self.task_outputs,
                    "raw": "Crew result",
                },
            )(),
            "Example project",
            mode,
            document_context,
        )

    def test_general_mode_generates_exactly_four_core_documents(self):
        files, archive_path, validation_report = self._write_package()
        try:
            generated_names = {Path(file_path).name for file_path in files}
            expected_names = {filename for filename, _ in GENERAL_SOFTWARE_DOCUMENTS}

            self.assertEqual(len(files), 4)
            self.assertEqual(generated_names, expected_names)
            with zipfile.ZipFile(archive_path) as archive:
                archive_names = set(archive.namelist())
                self.assertEqual(archive.testzip(), None)
                self.assertEqual(
                    archive_names,
                    expected_names | {"validation_report.json"},
                )
                validation_json = json.loads(
                    archive.read("validation_report.json").decode("utf-8")
                )
                plan = _read_docx(archive.read("Software_Test_Plan.docx"))
                protocol = _read_docx(archive.read("Software_Test_Protocol.docx"))
                matrix = _read_xlsx(archive.read("Requirements_Traceability_Matrix.xlsx"))
                summary = _read_docx(
                    archive.read("Software_Test_Summary_Report.docx")
                )
                self.assertIn(
                    "word/document.xml",
                    _zip_members(archive.read("Software_Test_Plan.docx")),
                )
                self.assertIn(
                    "xl/workbook.xml",
                    _zip_members(
                        archive.read("Requirements_Traceability_Matrix.xlsx")
                    ),
                )
                self.assertEqual(
                    [Path(path).suffix for path in files],
                    [".docx", ".docx", ".xlsx", ".docx"],
                )
                matrix_workbook = load_workbook(
                    BytesIO(archive.read("Requirements_Traceability_Matrix.xlsx")),
                    read_only=True,
                    data_only=True,
                )
                self.assertEqual(matrix_workbook.active["A5"].value, "Requirement ID")
                self.assertEqual(matrix_workbook.active["B5"].value, "Test Case ID")

            self.assertIn("Project scope and acceptance criteria", plan)
            self.assertIn("Test strategy and environment", plan)
            self.assertIn("Manual and automated testing approach", plan)
            self.assertIn("Test execution procedures", protocol)
            self.assertIn("TC-1: Verify inventory item creation.", protocol)
            self.assertNotIn("| REQ-F-001 | TC-1 | Pending |", protocol)
            self.assertIn(
                "REQ-F-001 | TC-1 | Pending | Not provided",
                matrix,
            )
            self.assertNotIn("TC-1: Verify inventory item creation.", matrix)
            self.assertIn("Pending / Not Executed", summary)
            self.assertIn("no actual test execution evidence", summary)
            self.assertIn("Defects:", summary)
            self.assertIn("Not measured", summary)
            self.assertIn("Conclusion:", summary)
            self.assertEqual(validation_json, validation_report)
            self.assertEqual(len(validation_report["checks"]), 9)
        finally:
            self._remove_package(files, archive_path)

    def test_unsupported_execution_claims_are_sanitized_before_export(self):
        outputs = list(self.task_outputs)
        outputs[6] = TaskOutput(
            "## Requirements Traceability Matrix\n"
            "| Requirement ID | Test Case ID | Verification Status | Evidence Reference |\n"
            "| REQ-F-001 | TC-1 | Pending | Not provided |\n\n"
            "## Detailed Test Cases\n"
            "### TC-1\n"
            "- **Linked Requirement ID:** REQ-F-001\n"
            "- **Test Steps:** Create an inventory item\n"
            "- **Expected Results:** Inventory item is created\n"
            "- **Actual Results:** Inventory item was created\n"
            "- **Test Status:** Passed\n"
            "- **Objective Evidence Reference:** run-999"
        )
        outputs[9] = TaskOutput("All tests passed. Coverage is 100%.")
        files, archive_path, report = self._write_package(task_outputs=outputs)
        try:
            protocol_path = next(
                path for path in files if Path(path).name == "Software_Test_Protocol.docx"
            )
            protocol = _read_docx(Path(protocol_path).read_bytes())
            summary_path = next(
                path for path in files
                if Path(path).name == "Software_Test_Summary_Report.docx"
            )
            summary = _read_docx(Path(summary_path).read_bytes())

            self.assertNotIn("Inventory item was created", protocol)
            self.assertIn("Actual Results: Not executed", protocol)
            self.assertNotIn("All tests passed", summary)
            self.assertNotIn("Coverage is 100%", summary)
            execution_check = next(
                check for check in report["checks"]
                if check["id"] == "execution_results_supported"
            )
            self.assertEqual(execution_check["status"], "REVIEW")
            self.assertIn("TC-1", execution_check["findings"][0])
            self.assertTrue(any("Unsupported execution summary" in item for item in execution_check["findings"]))
        finally:
            self._remove_package(files, archive_path)

    def test_medical_unsupported_execution_is_reported_before_export(self):
        outputs = list(self.task_outputs)
        outputs[4] = TaskOutput("RISK-1; RC-1")
        outputs[6] = TaskOutput(
            "## Requirements and Risk Traceability Matrix\n"
            "| Requirement ID | Risk Reference | Risk Control | Test Case ID | "
            "Verification Status | Evidence Reference |\n"
            "|---|---|---|---|---|---|\n"
            "| REQ-F-001 | RISK-1 | RC-1 | TC-1 | Not Executed | Not provided |\n\n"
            "## Verification Test Cases\n"
            "### TC-1\n"
            "- **Linked Requirement ID:** REQ-F-001\n"
            "- **Linked Risk Control ID:** RC-1\n"
            "- **Expected Results:** Inventory item is created\n"
            "- **Actual Results:** Inventory item was created\n"
            "- **Test Status:** Passed\n"
            "- **Objective Evidence Reference:** run-1"
        )
        files, archive_path, report = self._write_package(
            DocumentationMode.MEDICAL_DEVICE_SOFTWARE,
            task_outputs=outputs,
        )
        try:
            protocol_path = next(
                path
                for path in files
                if Path(path).name == "Software_Verification_Protocol.docx"
            )
            protocol = _read_docx(Path(protocol_path).read_bytes())
            execution_check = next(
                check
                for check in report["checks"]
                if check["id"] == "execution_results_supported"
            )

            self.assertNotIn("Inventory item was created", protocol)
            self.assertIn("Actual Results: Not executed", protocol)
            self.assertEqual(execution_check["status"], "REVIEW")
            self.assertTrue(execution_check["findings"])
        finally:
            self._remove_package(files, archive_path)

    def test_medical_device_mode_generates_four_verification_documents(self):
        medical_outputs = list(self.task_outputs)
        medical_outputs[6] = TaskOutput(
            "## Requirements and Risk Traceability Matrix\n"
            "| Requirement ID | Risk Reference | Risk Control | Test Case ID | "
            "Verification Status | Evidence Reference |\n"
            "| REQ-F-001 | RISK-1 | RC-1 | TC-1 | Pending / Not Verified | Not Provided |\n\n"
            "## Verification Test Cases\n"
            "### TC-1\n"
            "- **Linked Requirement ID:** REQ-F-001\n"
            "- **Linked Risk Control ID:** RC-1\n"
            "- **Preconditions:** Approved test environment available\n"
            "- **Test Steps:** Execute the documented verification steps\n"
            "- **Expected Results:** System meets REQ-F-001\n"
            "- **Actual Results:** Not provided\n"
            "- **Test Status:** Not Executed\n"
            "- **Objective Evidence Reference:** Not provided"
        )
        files, archive_path, validation_report = self._write_package(
            DocumentationMode.MEDICAL_DEVICE_SOFTWARE,
            task_outputs=medical_outputs,
        )
        try:
            generated_names = {Path(file_path).name for file_path in files}
            expected_names = {
                filename for filename, _ in MEDICAL_DEVICE_SOFTWARE_DOCUMENTS
            }

            self.assertEqual(len(files), 4)
            self.assertEqual(generated_names, expected_names)
            with zipfile.ZipFile(archive_path) as archive:
                self.assertEqual(
                    set(archive.namelist()),
                    expected_names | {"validation_report.json"},
                )
                report_from_archive = json.loads(
                    archive.read("validation_report.json").decode("utf-8")
                )
                plan = _read_docx(archive.read("Software_Test_Plan.docx"))
                protocol = _read_docx(
                    archive.read("Software_Verification_Protocol.docx")
                )
                matrix = _read_xlsx(
                    archive.read("Requirements_Risk_Traceability_Matrix.xlsx")
                )
                summary = _read_docx(
                    archive.read("Software_Verification_Summary_Report.docx")
                )
                self.assertIn(
                    "xl/workbook.xml",
                    _zip_members(
                        archive.read("Requirements_Risk_Traceability_Matrix.xlsx")
                    ),
                )
                medical_matrix = load_workbook(
                    BytesIO(
                        archive.read("Requirements_Risk_Traceability_Matrix.xlsx")
                    ),
                    read_only=True,
                    data_only=True,
                )
                self.assertEqual(medical_matrix.active["A5"].value, "Requirement ID")
                self.assertEqual(medical_matrix.active["C5"].value, "Risk Control")

            self.assertIn("Project scope and acceptance criteria", plan)
            for section in (
                "Intended Use and Scope",
                "Applicable Standards",
                "Software Safety Classification",
                "Risk-Based Testing Strategy",
                "Verification Approach",
                "Acceptance Criteria",
                "Review and Approval",
            ):
                self.assertIn(section, plan)
            self.assertIn("classification not determined", plan)
            self.assertIn("TC-1", protocol)
            for field in (
                "Linked Requirement ID",
                "Linked Risk Control ID",
                "Preconditions",
                "Test Steps",
                "Expected Results",
                "Actual Results",
                "Test Status",
                "Objective Evidence Reference",
            ):
                self.assertIn(field, protocol)
            self.assertIn(
                "REQ-F-001 | RISK-1 | RC-1 | TC-1 | Pending / Not Verified | Not provided",
                matrix,
            )
            self.assertNotIn("TC-1", summary.split("## Open Anomalies", 1)[0])
            self.assertIn("Pending / Not Executed", summary)
            self.assertIn("Risk-Control Verification", summary)
            self.assertIn("Review and Approval", summary)
            self.assertIn("Not provided | Pending", summary)
            self.assertNotIn("Executive conclusions", summary)
            self.assertEqual(validation_report["documentation_mode"], DocumentationMode.MEDICAL_DEVICE_SOFTWARE.value)
            self.assertEqual(report_from_archive, validation_report)
        finally:
            self._remove_package(files, archive_path)

    def test_medical_device_mode_marks_missing_content_as_pending(self):
        empty_outputs = [TaskOutput("") for _ in range(10)]
        files, archive_path, _ = self._write_package(
            DocumentationMode.MEDICAL_DEVICE_SOFTWARE,
            task_outputs=empty_outputs,
        )
        try:
            with zipfile.ZipFile(archive_path) as archive:
                plan = _read_docx(archive.read("Software_Test_Plan.docx"))
                protocol = _read_docx(
                    archive.read("Software_Verification_Protocol.docx")
                )
                matrix = _read_xlsx(
                    archive.read("Requirements_Risk_Traceability_Matrix.xlsx")
                )
                summary = _read_docx(
                    archive.read("Software_Verification_Summary_Report.docx")
                )

            self.assertIn("Not provided; classification not determined.", plan)
            self.assertIn("Test Case ID: Not provided", protocol)
            self.assertIn("Test Status: Not Executed", protocol)
            self.assertIn("Actual Results: Not executed", protocol)
            self.assertIn("Pending / Not Verified", matrix)
            self.assertIn("Pending / Not Executed", summary)
            self.assertNotIn("PASS", summary)
            self.assertNotIn("Approved by", summary)
        finally:
            self._remove_package(files, archive_path)

    def test_medical_summary_uses_supplied_execution_evidence(self):
        execution_context = (
            "Test Case ID: TC-1\n"
            "Actual Results: Inventory count updated\n"
            "Test Status: Passed\n"
            "Objective Evidence Reference: run-123/screenshot.png"
        )
        files, archive_path, _ = self._write_package(
            DocumentationMode.MEDICAL_DEVICE_SOFTWARE,
            document_context=execution_context,
        )
        try:
            with zipfile.ZipFile(archive_path) as archive:
                summary = _read_docx(
                    archive.read("Software_Verification_Summary_Report.docx")
                )

            self.assertIn("See supplied, test-case-linked execution evidence.", summary)
            self.assertIn("Project-Specific Summary", summary)
            self.assertIn("Executive conclusions", summary)
        finally:
            self._remove_package(files, archive_path)

    @staticmethod
    def _remove_package(files, archive_path):
        for file_path in files:
            Path(file_path).unlink(missing_ok=True)
        Path(files[0]).parent.rmdir()
        Path(archive_path).unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
